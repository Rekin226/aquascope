"""Study this area: a multi-gauge flood study over a set of stations or a bounding box.

One engine, three faces: :func:`study_area` is the function the Explorer's
browser worker, ``aquascope area-study`` and the MCP tool ``study_area`` all
call. It runs in CPython and in Pyodide with only numpy, scipy and pandas.

What it does, in order:

1. **Inventory**: the stations inside the area (or the list given) that record
   the variable, one per site, longest catalog record first, capped at
   :data:`MAX_SITES`.
2. **Read the Archive first**: each station's harvested daily file from the
   Hugging Face mirror. Only the stations the mirror does not hold are fetched
   live from the agency, and at most ``max_live`` of them (USGS allows about 50
   keyless requests an hour per address, so a box over a dense network would
   otherwise run out halfway). The rest are listed as skipped, with the reason.
3. **Per-site summary**: record span, mean, the annual maxima, Q100 from a GEV
   fitted by L-moments, and a Mann-Kendall trend on the annual maxima.
4. **Two regional methods**: field significance of the trends (counts up, down
   and none; Benjamini-Hochberg false discovery rate and the Walker test over
   the per-site p-values) and an index-flood regional frequency analysis
   (Hosking and Wallis: at-site growth curves, a pooled GEV growth curve from
   record-length weighted L-moment ratios, discordancy and a heterogeneity
   measure).

The result is plain JSON: a table, GeoJSON points carrying the per-site values
for colouring, the two regional blocks, notes and method citations. No model
sits between the numbers and the result.
"""

from __future__ import annotations

import io
import logging
import math
from collections.abc import Callable
from datetime import date
from typing import Any

import numpy as np
import pandas as pd

logger = logging.getLogger(__name__)

#: At most this many stations are studied in one area (archive reads included).
MAX_SITES = 60
#: At most this many stations are fetched live from an agency (the rest must come from the Archive).
MAX_LIVE_FETCHES = 25
#: Annual maxima needed for a site's trend test (the same floor as the single-station study).
MIN_YEARS_TREND = 8
#: Annual maxima needed for a site's flood frequency and a place in the regional pool.
MIN_YEARS_FFA = 10
#: Significance level for per-site trends and for field significance.
ALPHA = 0.05
#: Return periods the regional growth curve reports.
RETURN_PERIODS = [2, 5, 10, 25, 50, 100]
#: Regions simulated for the heterogeneity measure.
N_SIMULATIONS = 500

#: Critical values of the discordancy measure D (Hosking and Wallis 1997, Table 3.1), by number of sites.
DISCORDANCY_CRITICAL = {5: 1.333, 6: 1.648, 7: 1.917, 8: 2.140, 9: 2.329, 10: 2.491, 11: 2.632, 12: 2.757,
                        13: 2.869, 14: 2.971}

METHODS: dict[str, dict[str, str]] = {
    "amax_trend": {
        "name": "Mann-Kendall trend on annual maxima",
        "text": "Non-parametric Mann-Kendall test with Sen's slope on each site's annual maximum daily flow "
        "(calendar years with at least 80 % daily coverage).",
        "citation": "Mann, H. B. (1945). Nonparametric tests against trend. Econometrica, 13, 245-259; "
        "Sen, P. K. (1968). J. Am. Stat. Assoc., 63, 1379-1389.",
    },
    "field_significance": {
        "name": "Field significance (Benjamini-Hochberg FDR and Walker test)",
        "text": "Per-site trend p-values tested together: the false discovery rate procedure flags which sites "
        "stay significant, and the Walker test asks whether the smallest p-value is smaller than chance allows "
        "for this many sites. Both assume independent sites; nearby gauges are correlated, which makes a "
        "field look more significant than it is.",
        "citation": "Benjamini, Y., & Hochberg, Y. (1995). Controlling the false discovery rate. J. R. Stat. Soc. "
        "B, 57(1), 289-300; Wilks, D. S. (2006). On \"field significance\" and the false discovery rate. "
        "J. Appl. Meteorol. Climatol., 45, 1181-1189.",
    },
    "index_flood": {
        "name": "Index-flood regional frequency analysis (L-moments)",
        "text": "Each site's annual maxima scaled by its mean (the index flood); regional L-moment ratios "
        "weighted by record length; a GEV growth curve fitted to them; discordancy D and heterogeneity H "
        "(simulated from the regional GEV rather than the kappa distribution of the textbook).",
        "citation": "Dalrymple, T. (1960). Flood-frequency analyses. USGS Water-Supply Paper 1543-A; "
        "Hosking, J. R. M., & Wallis, J. R. (1997). Regional Frequency Analysis: An Approach Based on "
        "L-Moments. Cambridge University Press.",
    },
    "gev_lmoments": {
        "name": "GEV fitted by L-moments",
        "text": "Each site's annual maxima fitted to a GEV with L-moment estimators; Q100 from the fitted "
        "quantile function.",
        "citation": "Hosking, J. R. M. (1990). L-moments. J. R. Stat. Soc. B, 52(1), 105-124.",
    },
}

#: Column order of the result table (and of the CSV and XLSX downloads).
TABLE_COLUMNS = [
    "source", "station_id", "name", "latitude", "longitude", "status", "data_from", "start", "end",
    "record_years", "n_amax", "mean", "unit", "area_km2", "index_flood", "q100", "q100_per_km2",
    "regional_q100", "trend", "trend_p", "sens_slope_per_year", "fdr_significant", "discordancy", "note",
]

ProgressFn = Callable[[dict[str, Any]], None]


# ── small helpers ───────────────────────────────────────────────────────────


def _num(x: Any, digits: int = 4) -> Any:
    """JSON-safe number: NaN and inf become None; floats keep ``digits`` significant decimals."""
    if x is None:
        return None
    try:
        v = float(x)
    except (TypeError, ValueError):
        return None
    if math.isnan(v) or math.isinf(v):
        return None
    if v == 0:
        return 0.0
    mag = max(0, digits - 1 - int(math.floor(math.log10(abs(v)))))
    return round(v, mag)


def _key(source: str, station_id: str) -> str:
    return f"{source}/{station_id}"


def _span_days(row: dict[str, Any], today: date) -> int:
    try:
        start = date.fromisoformat(str(row.get("period_start"))[:10])
    except (TypeError, ValueError):
        return -1
    try:
        end = date.fromisoformat(str(row.get("period_end"))[:10]) if row.get("period_end") else today
    except (TypeError, ValueError):
        end = today
    return max(-1, (end - start).days)


def _normalise_station(item: Any) -> dict[str, Any] | None:
    """A station as ``{"source", "station_id", ...}`` from a dict (catalog row or page row) or a "source/id"."""
    if isinstance(item, str):
        if "/" not in item:
            return None
        source, station_id = item.split("/", 1)
        return {"source": source.strip(), "station_id": station_id.strip()}
    if not isinstance(item, dict) or not item.get("source") or not item.get("station_id"):
        return None
    row = dict(item)
    if "latitude" not in row and "lat" in row:
        row["latitude"] = row.get("lat")
    if "longitude" not in row and "lon" in row:
        row["longitude"] = row.get("lon")
    extra = row.get("extra") if isinstance(row.get("extra"), dict) else {}
    if row.get("area_km2") in (None, "") and extra.get("catchment_area_km2"):
        row["area_km2"] = extra.get("catchment_area_km2")
    row["source"], row["station_id"] = str(row["source"]), str(row["station_id"])
    return row


# ── 1. inventory ────────────────────────────────────────────────────────────


def inventory(
    stations: list[Any] | None = None,
    *,
    bbox: tuple[float, float, float, float] | list[float] | None = None,
    variable: str = "discharge",
    rows: list[dict[str, Any]] | None = None,
    max_sites: int = MAX_SITES,
) -> dict[str, Any]:
    """The stations a study of this area would use.

    ``stations`` is a list of catalog-like dicts (``source``, ``station_id``,
    optional ``name``, ``latitude``/``lat``, ``longitude``/``lon``,
    ``variables``, ``period_start``, ``period_end``, ``area_km2``) or
    ``"source/station_id"`` strings. Without it, ``bbox`` = (west, south, east,
    north) selects from ``rows`` (default: the published catalog). Stations whose
    catalog row says they do not record ``variable`` are left out; one station
    per site is kept (the longest record); the longest ``max_sites`` are studied.

    Returns ``{"sites": [...], "n_in_area", "n_without_variable", "n_over_cap", "notes"}``.
    """
    notes: list[str] = []
    if stations is None:
        if bbox is None:
            raise ValueError("give stations or a bbox")
        from aquascope.schemas.station import in_bbox

        if rows is None:
            from aquascope.archive.catalog import load_stations

            rows = load_stations()
        box = tuple(float(x) for x in bbox)
        if len(box) != 4:
            raise ValueError("bbox is (west, south, east, north)")
        west, south, east, north = box
        picked = []
        for r in rows:
            lat, lon = r.get("latitude"), r.get("longitude")
            if lat is None or lon is None:
                continue
            if west <= east:
                inside = in_bbox(float(lat), float(lon), box)  # type: ignore[arg-type]
            else:  # across the antimeridian
                inside = south <= float(lat) <= north and (float(lon) >= west or float(lon) <= east)
            if inside:
                picked.append(r)
        stations = picked
    candidates = [s for s in (_normalise_station(x) for x in stations) if s is not None]
    n_in_area = len(candidates)
    with_var = [s for s in candidates if not s.get("variables") or variable in (s.get("variables") or [])]
    n_without = n_in_area - len(with_var)
    # One station per site: several records at one place would count the same floods twice.
    today = date.today()
    by_site: dict[tuple[str, str], dict[str, Any]] = {}
    for s in with_var:
        k = (s["source"], str(s.get("site_id") or s["station_id"]))
        if k not in by_site or _span_days(s, today) > _span_days(by_site[k], today):
            by_site[k] = s
    n_dupes = len(with_var) - len(by_site)
    sites = sorted(by_site.values(), key=lambda s: (-_span_days(s, today), s["source"], s["station_id"]))
    n_over = max(0, len(sites) - int(max_sites))
    if n_without:
        notes.append(f"{n_without} station(s) in the area do not record {variable.replace('_', ' ')}.")
    if n_dupes:
        notes.append(f"{n_dupes} duplicate record(s) at the same site were merged, keeping the longest.")
    if n_over:
        notes.append(f"The area holds {len(sites)} {variable.replace('_', ' ')} sites; the {max_sites} with the "
                     f"longest catalog records are studied and {n_over} are left out.")
    return {"sites": sites[: int(max_sites)], "dropped": sites[int(max_sites):], "n_in_area": n_in_area,
            "n_without_variable": n_without, "n_over_cap": n_over, "notes": notes}


# ── 2. fetching: the Archive first, then a capped number of live calls ──────


def _archive_reader(source: str, station_id: str, variable: str) -> pd.Series | None:
    from aquascope.archive.observations import fetch_archived_series, harvestable_variables

    if variable not in harvestable_variables(source):
        return None
    return fetch_archived_series(source, station_id, variable)


def _live_reader(source: str, station_id: str, variable: str, period_start: Any = None) -> pd.Series | None:
    from aquascope.explore import fetch_series

    got = fetch_series(source, station_id, variable=variable, prefer_archive=False, period_start=period_start)
    return got.get("series")


def gather_series(
    sites: list[dict[str, Any]],
    *,
    variable: str = "discharge",
    max_live: int = MAX_LIVE_FETCHES,
    archive_reader: Callable[..., pd.Series | None] | None = None,
    live_reader: Callable[..., pd.Series | None] | None = None,
    on_progress: ProgressFn | None = None,
) -> tuple[dict[str, pd.Series], dict[str, dict[str, Any]]]:
    """Read each site's record: the Archive mirror first, then at most ``max_live`` live agency calls.

    Returns ``(series by key, status by key)``; a status is ``{"status": "studied" | "skipped" | "failed",
    "data_from": "archive" | "live" | None, "note": str}``.
    """
    from aquascope.registry import SOURCES

    archive_reader = archive_reader or _archive_reader
    live_reader = live_reader or _live_reader
    series: dict[str, pd.Series] = {}
    status: dict[str, dict[str, Any]] = {}
    live_left = max(0, int(max_live))
    total = len(sites)
    for i, site in enumerate(sites):
        src, sid = site["source"], site["station_id"]
        key = _key(src, sid)
        if on_progress:
            on_progress({"phase": "fetch", "done": i, "total": total, "site": key, "name": site.get("name")})
        s = None
        try:
            s = archive_reader(src, sid, variable)
        except Exception as exc:  # noqa: BLE001 - the Archive is a fast path, never a hard dependency
            logger.info("archive read failed for %s: %s", key, exc)
        if s is not None and not s.dropna().empty:
            series[key] = s.dropna()
            status[key] = {"status": "studied", "data_from": "archive", "note": ""}
            continue
        if src not in SOURCES:
            status[key] = {"status": "failed", "data_from": None, "note": "unknown source"}
            continue
        if live_left <= 0:
            status[key] = {"status": "skipped", "data_from": None,
                           "note": f"not in the Archive and the live-fetch cap ({max_live}) was reached"}
            continue
        live_left -= 1
        try:
            s = live_reader(src, sid, variable, site.get("period_start"))
        except Exception as exc:  # noqa: BLE001 - one agency failing must not sink the area
            msg = str(exc).splitlines()[0][:160] if str(exc) else type(exc).__name__
            status[key] = {"status": "failed", "data_from": "live", "note": msg}
            continue
        if s is None or s.dropna().empty:
            status[key] = {"status": "failed", "data_from": "live", "note": "the agency returned no observations"}
            continue
        series[key] = s.dropna()
        status[key] = {"status": "studied", "data_from": "live", "note": ""}
    if on_progress:
        on_progress({"phase": "fetch", "done": total, "total": total, "site": None})
    return series, status


# ── 3. per-site summary ─────────────────────────────────────────────────────


def annual_maxima(s: pd.Series) -> pd.Series:
    """Calendar-year maxima of daily means, keeping years with at least ~80 % daily coverage."""
    from aquascope.explore import _annual_max

    return _annual_max(s)


def site_summary(s: pd.Series, *, area_km2: float | None = None, alpha: float = ALPHA) -> dict[str, Any]:
    """Record span, mean, annual maxima, Q100 (GEV L-moments) and the annual-maxima trend for one series.

    Pure: no fetching. ``trend_p`` keeps full precision (the regional tests need it).
    """
    s = s.dropna()
    out: dict[str, Any] = {
        "start": s.index.min().date().isoformat() if len(s) else None,
        "end": s.index.max().date().isoformat() if len(s) else None,
        "record_years": _num((s.index.max() - s.index.min()).days / 365.25, 3) if len(s) > 1 else 0.0,
        "mean": _num(s.mean()) if len(s) else None,
        "n_amax": 0, "index_flood": None, "q100": None, "q100_per_km2": None,
        "trend": "untested", "trend_p": None, "trend_tau": None, "sens_slope_per_year": None,
        "area_km2": _num(area_km2) if area_km2 else None,
    }
    if not len(s):
        return out
    am = annual_maxima(s)
    out["n_amax"] = int(len(am))
    out["amax"] = {"year": [int(y) for y in am.index.year], "v": [float(v) for v in am.values]}
    if len(am):
        out["index_flood"] = _num(am.mean())
    if len(am) >= MIN_YEARS_FFA:
        from aquascope.hydrology.flood_frequency import fit_gev_lmoments

        try:
            g = fit_gev_lmoments(am, return_periods=[100])
            out["q100"] = _num(g.return_periods[100])
            if area_km2 and float(area_km2) > 0 and out["q100"] is not None:
                out["q100_per_km2"] = _num(out["q100"] / float(area_km2))
        except Exception as exc:  # noqa: BLE001
            logger.info("GEV fit failed: %s", exc)
    if len(am) >= MIN_YEARS_TREND:
        from aquascope.analysis.trends import mann_kendall, sens_slope

        try:
            mk = mann_kendall(am.values)
            p = float(mk.p_value)
            tau = float(mk.tau)
            out["trend_p"] = p
            out["trend_tau"] = _num(tau)
            out["sens_slope_per_year"] = _num(sens_slope(am.values).slope)
            out["trend"] = ("up" if tau > 0 else "down") if p < alpha else "none"
        except Exception as exc:  # noqa: BLE001
            logger.info("trend skipped: %s", exc)
    return out


# ── 4a. field significance ──────────────────────────────────────────────────


def benjamini_hochberg(pvalues: list[float], q: float = ALPHA) -> list[bool]:
    """Which hypotheses the Benjamini-Hochberg procedure rejects at false discovery rate ``q``."""
    p = np.asarray(pvalues, dtype=float)
    n = len(p)
    if n == 0:
        return []
    order = np.argsort(p)
    thresholds = q * np.arange(1, n + 1) / n
    below = p[order] <= thresholds
    reject = np.zeros(n, dtype=bool)
    if below.any():
        k = int(np.max(np.nonzero(below)[0]))
        reject[order[: k + 1]] = True
    return [bool(x) for x in reject]


def field_significance(sites: dict[str, dict[str, Any]], *, alpha: float = ALPHA) -> dict[str, Any]:
    """Are the per-site annual-maxima trends, taken together, more than chance?

    ``sites`` maps a key to a :func:`site_summary`. Returns the counts up, down and none (at ``alpha`` per
    site), the sites the Benjamini-Hochberg procedure keeps at FDR ``alpha``, the Walker test, a binomial
    count test, a one-line verdict and the independence caveat.
    """
    tested = {k: v for k, v in sites.items() if v.get("trend_p") is not None}
    keys = list(tested)
    n = len(keys)
    up = sum(1 for k in keys if tested[k]["trend"] == "up")
    down = sum(1 for k in keys if tested[k]["trend"] == "down")
    out: dict[str, Any] = {
        "alpha": alpha, "n_tested": n, "n_untested": len(sites) - n,
        "counts": {"up": up, "down": down, "none": n - up - down},
        "expected_by_chance": _num(alpha * n, 3),
        "fdr": {"q": alpha, "significant": [], "n_significant": 0},
        "walker": None, "binomial_p": None, "field_significant": None,
        "caveat": "Both tests treat the sites as independent. Nearby gauges share floods, so the real chance of "
                  "a field this strong is higher than shown; a block bootstrap that keeps the years together "
                  "would account for it.",
    }
    if n == 0:
        out["verdict"] = f"No site has the {MIN_YEARS_TREND} complete years of annual maxima a trend test needs."
        return out
    pvals = [float(tested[k]["trend_p"]) for k in keys]
    rejected = benjamini_hochberg(pvals, q=alpha)
    sig = [k for k, r in zip(keys, rejected) if r]
    out["fdr"] = {"q": alpha, "significant": sig, "n_significant": len(sig)}
    p_min = min(pvals)
    walker_crit = 1.0 - (1.0 - alpha) ** (1.0 / n)
    out["walker"] = {"p_min": _num(p_min), "critical": _num(walker_crit), "significant": bool(p_min <= walker_crit)}
    from scipy.stats import binomtest

    n_sig = up + down
    out["binomial_p"] = _num(binomtest(n_sig, n, alpha, alternative="greater").pvalue)
    out["field_significant"] = bool(sig) or out["walker"]["significant"]
    if n < 3:
        out["verdict"] = (f"Only {n} site(s) could be tested, too few for a regional statement; "
                          f"{up} up, {down} down at p < {alpha:g}.")
    elif out["field_significant"]:
        lean = "upward" if up > down else "downward" if down > up else "mixed"
        out["verdict"] = (f"{up} up, {down} down, {n - up - down} none of {n} sites (p < {alpha:g}); "
                          f"{len(sig)} stay significant after the false discovery rate, so the {lean} signal is "
                          f"more than chance if the sites were independent.")
    else:
        out["verdict"] = (f"{up} up, {down} down, {n - up - down} none of {n} sites (p < {alpha:g}); none survive "
                          f"the false discovery rate and the Walker test is not significant, so the field shows "
                          f"no trend beyond chance.")
    return out


# ── 4b. index-flood regional frequency analysis ─────────────────────────────


def _lmoment_ratios(x: np.ndarray) -> tuple[float, float, float, float]:
    """Sample (l1, t, t3, t4) by unbiased probability-weighted moments (vectorised)."""
    x = np.sort(np.asarray(x, dtype=float))
    n = len(x)
    i = np.arange(n, dtype=float)
    b0 = x.mean()
    b1 = np.sum(i / (n - 1) * x) / n
    b2 = np.sum(i * (i - 1) / ((n - 1) * (n - 2)) * x) / n
    b3 = np.sum(i * (i - 1) * (i - 2) / ((n - 1) * (n - 2) * (n - 3)) * x) / n
    l1, l2 = b0, 2 * b1 - b0
    l3 = 6 * b2 - 6 * b1 + b0
    l4 = 20 * b3 - 30 * b2 + 12 * b1 - b0
    if l2 <= 0 or l1 == 0:
        return float(l1), float("nan"), float("nan"), float("nan")
    return float(l1), float(l2 / l1), float(l3 / l2), float(l4 / l2)


def _lcv_rows(x: np.ndarray) -> np.ndarray:
    """Sample L-CV of every row of ``x`` (rows are samples), vectorised for the heterogeneity simulation."""
    x = np.sort(x, axis=1)
    n = x.shape[1]
    i = np.arange(n, dtype=float)
    b0 = x.mean(axis=1)
    b1 = (x * (i / (n - 1))).sum(axis=1) / n
    with np.errstate(divide="ignore", invalid="ignore"):
        return (2 * b1 - b0) / b0


def _gev_from_ratios(t: float, t3: float) -> tuple[float, float, float]:
    """GEV (shape k, location, scale) with mean 1, L-CV ``t`` and L-skewness ``t3`` (Hosking 1990).

    ``k`` follows scipy's ``genextreme`` sign convention.
    """
    c = 2.0 / (3.0 + t3) - math.log(2.0) / math.log(3.0)
    k = 7.8590 * c + 2.9554 * c * c
    l2 = t  # mean 1, so L2 = L-CV
    if abs(k) < 1e-8:
        alpha = l2 / math.log(2.0)
        xi = 1.0 - alpha * 0.5772156649
    else:
        g = math.gamma(1.0 + k)
        alpha = l2 * k / (g * (1.0 - 2.0 ** (-k)))
        xi = 1.0 - alpha * (1.0 - g) / k
    return k, xi, alpha


def _discordancy(u: np.ndarray) -> list[float]:
    """Hosking-Wallis discordancy D_i over rows u_i = (t, t3, t4)."""
    n = len(u)
    d = u - u.mean(axis=0)
    a = d.T @ d
    a_inv = np.linalg.pinv(a)
    return [float(n / 3.0 * (row @ a_inv @ row)) for row in d]


def index_flood_rfa(
    amax: dict[str, np.ndarray | pd.Series | list[float]],
    *,
    return_periods: list[float] | None = None,
    n_sim: int = N_SIMULATIONS,
    seed: int = 0,
) -> dict[str, Any]:
    """Index-flood regional frequency analysis over the sites' annual maxima.

    Sites with fewer than :data:`MIN_YEARS_FFA` maxima are left out. Returns the regional L-moment ratios, the
    GEV growth curve, per-site index floods, at-site and regional quantiles, discordancy and the heterogeneity
    measure H with its class, plus notes. ``seed`` makes H reproducible.
    """
    from scipy.stats import genextreme

    rps = [float(x) for x in (return_periods or RETURN_PERIODS)]
    pool = {k: np.asarray(v, dtype=float) for k, v in amax.items()}
    pool = {k: v[np.isfinite(v)] for k, v in pool.items()}
    left_out = sorted(k for k, v in pool.items() if len(v) < MIN_YEARS_FFA)
    pool = {k: v for k, v in pool.items() if len(v) >= MIN_YEARS_FFA and v.mean() > 0}
    out: dict[str, Any] = {"n_sites": len(pool), "left_out": left_out, "return_periods": rps, "notes": []}
    if len(pool) < 2:
        out["error"] = (f"A regional growth curve needs at least 2 sites with {MIN_YEARS_FFA} years of annual "
                        f"maxima; this area has {len(pool)}.")
        return out
    keys = list(pool)
    ratios = {k: _lmoment_ratios(pool[k]) for k in keys}
    n_i = np.array([len(pool[k]) for k in keys], dtype=float)
    t_i = np.array([ratios[k][1] for k in keys])
    t3_i = np.array([ratios[k][2] for k in keys])
    t4_i = np.array([ratios[k][3] for k in keys])
    w = n_i / n_i.sum()
    t_r, t3_r, t4_r = float(w @ t_i), float(w @ t3_i), float(w @ t4_i)
    k_r, xi_r, a_r = _gev_from_ratios(t_r, t3_r)
    growth = {f"{rp:g}": _num(genextreme.ppf(1 - 1 / rp, k_r, loc=xi_r, scale=a_r)) for rp in rps}
    out["regional"] = {"l_cv": _num(t_r), "l_skew": _num(t3_r), "l_kurt": _num(t4_r),
                       "gev": {"shape": _num(k_r), "location": _num(xi_r), "scale": _num(a_r)},
                       "growth_curve": growth, "n_years": int(n_i.sum())}

    sites: dict[str, dict[str, Any]] = {}
    for j, k in enumerate(keys):
        index = float(ratios[k][0])
        kk, xx, aa = _gev_from_ratios(float(t_i[j]), float(t3_i[j]))
        at_site = {f"{rp:g}": _num(index * genextreme.ppf(1 - 1 / rp, kk, loc=xx, scale=aa)) for rp in rps}
        regional = {f"{rp:g}": _num(index * (growth[f"{rp:g}"] or float("nan"))) for rp in rps}
        sites[k] = {"n_years": int(n_i[j]), "index_flood": _num(index), "l_cv": _num(t_i[j]),
                    "l_skew": _num(t3_i[j]), "l_kurt": _num(t4_i[j]),
                    "growth_curve": {f"{rp:g}": _num(genextreme.ppf(1 - 1 / rp, kk, loc=xx, scale=aa))
                                     for rp in rps},
                    "q_at_site": at_site, "q_regional": regional}

    # Discordancy (meaningful from five sites up).
    n = len(keys)
    if n >= 5:
        d = _discordancy(np.column_stack([t_i, t3_i, t4_i]))
        crit = DISCORDANCY_CRITICAL.get(n, 3.0)
        for j, k in enumerate(keys):
            sites[k]["discordancy"] = _num(d[j], 3)
            sites[k]["discordant"] = bool(d[j] >= crit)
        flagged = [k for k in keys if sites[k]["discordant"]]
        out["discordancy"] = {"critical": crit, "discordant": flagged}
        if flagged:
            out["notes"].append(f"{len(flagged)} site(s) have unusual L-moment ratios (D >= {crit:g}): "
                                f"{', '.join(flagged)}. Check their data before trusting the pool.")
    else:
        out["discordancy"] = {"critical": None, "discordant": []}
        out["notes"].append("Discordancy needs at least 5 sites.")

    # Heterogeneity H1 on the L-CV (Hosking and Wallis 1997, section 4.3.3), simulated from the regional GEV.
    v_obs = math.sqrt(float(w @ (t_i - t_r) ** 2))
    rng = np.random.default_rng(seed)
    t_sim = np.column_stack([
        _lcv_rows(genextreme.ppf(rng.random((n_sim, int(n_i[j]))), k_r, loc=xi_r, scale=a_r)) for j in range(n)
    ])  # (n_sim, n): each simulated region's at-site L-CVs
    tr_sim = t_sim @ w
    v_sim = np.sqrt(((t_sim - tr_sim[:, None]) ** 2) @ w)
    mu, sd = float(v_sim.mean()), float(v_sim.std(ddof=1))
    h = (v_obs - mu) / sd if sd > 0 else float("nan")
    if not math.isfinite(h):
        klass = "unknown"
    elif h < 1:
        klass = "acceptably homogeneous"
    elif h < 2:
        klass = "possibly heterogeneous"
    else:
        klass = "definitely heterogeneous"
    out["heterogeneity"] = {"H": _num(h, 3), "V": _num(v_obs), "class": klass, "n_sim": n_sim,
                            "simulated_from": "regional GEV (the textbook uses a four-parameter kappa)"}
    out["sites"] = sites
    if n < 5:
        out["notes"].append(f"Only {n} sites in the pool; Hosking and Wallis suggest at least 5 to 20.")
    if klass != "acceptably homogeneous":
        out["notes"].append(f"H = {h:.2f} ({klass}): one growth curve may not fit the whole area; "
                            "split it by size or climate before using the regional quantiles.")
    return out


# ── the study ───────────────────────────────────────────────────────────────


def _headline(summary: dict[str, Any], fs: dict[str, Any], rfa: dict[str, Any]) -> str:
    parts = [f"{summary['n_studied']} of {summary['n_selected']} gauges studied "
             f"({summary['n_archive']} from the Archive, {summary['n_live']} live)."]
    if fs.get("n_tested"):
        parts.append(fs["verdict"])
    if rfa.get("regional"):
        het = rfa.get("heterogeneity") or {}
        g100 = (rfa["regional"].get("growth_curve") or {}).get("100")
        times = f"Q100 is {g100:g} times the index flood" if g100 is not None else "no Q100 growth factor"
        parts.append(f"Regional growth curve from {rfa['n_sites']} sites: {times}; "
                     f"H = {het.get('H')} ({het.get('class')}).")
    elif rfa.get("error"):
        parts.append(rfa["error"])
    return " ".join(parts)


def analyse_sites(
    sites: list[dict[str, Any]],
    series: dict[str, pd.Series],
    status: dict[str, dict[str, Any]],
    *,
    variable: str = "discharge",
    unit: str = "m3/s",
    alpha: float = ALPHA,
    n_sim: int = N_SIMULATIONS,
    on_progress: ProgressFn | None = None,
) -> dict[str, Any]:
    """Per-site summaries plus the two regional methods over series already in hand. Pure (no fetching)."""
    summaries: dict[str, dict[str, Any]] = {}
    total = len(series)
    for i, (key, s) in enumerate(series.items()):
        if on_progress:
            on_progress({"phase": "analyse", "done": i, "total": total, "site": key})
        meta = next((x for x in sites if _key(x["source"], x["station_id"]) == key), {})
        area = meta.get("area_km2")
        try:
            area = float(area) if area not in (None, "") else None
        except (TypeError, ValueError):
            area = None
        summaries[key] = site_summary(s, area_km2=area, alpha=alpha)

    if on_progress:
        on_progress({"phase": "regional", "done": 0, "total": 2, "site": None})
    fs = field_significance(summaries, alpha=alpha)
    rfa = index_flood_rfa({k: np.asarray(v.get("amax", {}).get("v") or [], dtype=float)
                           for k, v in summaries.items()}, n_sim=n_sim)
    fdr_sig = set(fs["fdr"]["significant"])
    rfa_sites = rfa.get("sites") or {}

    rows: list[dict[str, Any]] = []
    features: list[dict[str, Any]] = []
    for site in sites:
        key = _key(site["source"], site["station_id"])
        st = status.get(key, {"status": "skipped", "data_from": None, "note": ""})
        summ = summaries.get(key, {})
        reg = rfa_sites.get(key, {})
        row = {
            "key": key, "source": site["source"], "station_id": site["station_id"], "name": site.get("name") or "",
            "latitude": _num(site.get("latitude"), 7), "longitude": _num(site.get("longitude"), 7),
            "status": st["status"], "data_from": st.get("data_from"), "note": st.get("note") or "",
            "start": summ.get("start"), "end": summ.get("end"), "record_years": summ.get("record_years"),
            "n_amax": summ.get("n_amax"), "mean": summ.get("mean"), "unit": unit if summ else None,
            "area_km2": summ.get("area_km2"), "index_flood": summ.get("index_flood"), "q100": summ.get("q100"),
            "q100_per_km2": summ.get("q100_per_km2"),
            "regional_q100": (reg.get("q_regional") or {}).get("100"),
            "trend": summ.get("trend") if summ else None,
            "trend_p": _num(summ.get("trend_p")) if summ.get("trend_p") is not None else None,
            "trend_tau": summ.get("trend_tau"), "sens_slope_per_year": summ.get("sens_slope_per_year"),
            "fdr_significant": (key in fdr_sig) if summ.get("trend_p") is not None else None,
            "discordancy": reg.get("discordancy"), "discordant": reg.get("discordant"),
        }
        if summ and summ.get("n_amax", 0) < MIN_YEARS_TREND and not row["note"]:
            row["note"] = f"{summ.get('n_amax', 0)} complete years of annual maxima; too short for a trend"
        rows.append(row)
        lat, lon = row["latitude"], row["longitude"]
        if lat is not None and lon is not None:
            props = {k: v for k, v in row.items() if k not in ("latitude", "longitude")}
            features.append({"type": "Feature", "geometry": {"type": "Point", "coordinates": [lon, lat]},
                             "properties": props})

    n_studied = sum(1 for r in rows if r["status"] == "studied")
    summary = {
        "n_selected": len(sites), "n_studied": n_studied,
        "n_archive": sum(1 for r in rows if r["data_from"] == "archive" and r["status"] == "studied"),
        "n_live": sum(1 for r in rows if r["data_from"] == "live" and r["status"] == "studied"),
        "n_skipped": sum(1 for r in rows if r["status"] == "skipped"),
        "n_failed": sum(1 for r in rows if r["status"] == "failed"),
    }
    if on_progress:
        on_progress({"phase": "regional", "done": 2, "total": 2, "site": None})
    return {
        "variable": variable, "unit": unit, "alpha": alpha,
        "summary": summary,
        "headline": _headline(summary, fs, rfa),
        "sites": rows,
        "table": {"columns": TABLE_COLUMNS, "rows": [[r.get(c) for c in TABLE_COLUMNS] for r in rows]},
        "geojson": {"type": "FeatureCollection", "features": features},
        "field_significance": fs,
        "regional_frequency": rfa,
        "methods": [METHODS[m] for m in ("amax_trend", "field_significance", "gev_lmoments", "index_flood")],
    }


def study_area(
    stations: list[Any] | None = None,
    *,
    bbox: tuple[float, float, float, float] | list[float] | None = None,
    question: str | None = None,
    variable: str = "discharge",
    rows: list[dict[str, Any]] | None = None,
    max_sites: int = MAX_SITES,
    max_live: int = MAX_LIVE_FETCHES,
    alpha: float = ALPHA,
    n_sim: int = N_SIMULATIONS,
    archive_reader: Callable[..., pd.Series | None] | None = None,
    live_reader: Callable[..., pd.Series | None] | None = None,
    on_progress: ProgressFn | None = None,
) -> dict[str, Any]:
    """Study the gauges of an area together: per-site flood summaries plus two regional methods.

    ``stations`` (dicts or ``"source/id"`` strings) or ``bbox`` = (west, south, east, north) pick the gauges;
    see :func:`inventory`. The Archive is read first and at most ``max_live`` stations are fetched live.
    ``question`` is carried through to the result as the study's title; it does not change the computation.
    ``on_progress`` receives ``{"phase", "done", "total", "site"}`` dicts as the study goes.

    Returns a JSON-safe dict: ``summary``, ``headline``, ``sites`` (one row per selected gauge), ``table``,
    ``geojson`` (points with the per-site values), ``field_significance``, ``regional_frequency``,
    ``skipped``, ``notes`` and ``methods``.
    """
    from aquascope.archive.observations import ARCHIVE_UNITS

    inv = inventory(stations, bbox=bbox, variable=variable, rows=rows, max_sites=max_sites)
    sites = inv["sites"]
    if on_progress:
        on_progress({"phase": "inventory", "done": len(sites), "total": len(sites), "site": None})
    series, status = gather_series(sites, variable=variable, max_live=max_live, archive_reader=archive_reader,
                                   live_reader=live_reader, on_progress=on_progress)
    result = analyse_sites(sites, series, status, variable=variable, unit=ARCHIVE_UNITS.get(variable, ""),
                           alpha=alpha, n_sim=n_sim, on_progress=on_progress)
    notes = list(inv["notes"])
    n_skipped = result["summary"]["n_skipped"]
    if n_skipped:
        notes.append(f"{n_skipped} gauge(s) were not in the Archive and were skipped once {max_live} live "
                     "fetches had been used (keyless agency limits). Run again later or with fewer gauges.")
    if result["summary"]["n_failed"]:
        notes.append(f"{result['summary']['n_failed']} gauge(s) could not be read; see the note column.")
    result.update({
        "question": question or f"Flood study of {len(sites)} gauges",
        "bbox": [float(x) for x in bbox] if bbox is not None else None,
        "n_in_area": inv["n_in_area"],
        "caps": {"max_sites": max_sites, "max_live": max_live},
        "skipped": [{"key": r["key"], "reason": r["note"]} for r in result["sites"] if r["status"] != "studied"]
        + [{"key": _key(s["source"], s["station_id"]), "reason": "over the site cap"} for s in inv["dropped"]],
        "notes": notes + list(result["regional_frequency"].get("notes") or []),
    })
    return result


def apply_areas(result: dict[str, Any], areas: dict[str, float | None]) -> dict[str, Any]:
    """Fill in catchment areas the caller learnt after the study (the Explorer reads them from the Archive's
    station_catchments table on the page) and recompute Q100 per km2 in the rows, table and GeoJSON.

    ``areas`` maps ``"source/station_id"`` to km2. A site keeps an area it already had. Returns ``result``.
    """
    by_key: dict[str, dict[str, Any]] = {}
    for r in result.get("sites") or []:
        a = areas.get(r["key"])
        try:
            a = float(a) if a is not None else None
        except (TypeError, ValueError):
            a = None
        if r.get("area_km2") is None and a and a > 0:
            r["area_km2"] = _num(a)
        if r.get("area_km2") and r.get("q100") is not None:
            r["q100_per_km2"] = _num(float(r["q100"]) / float(r["area_km2"]))
        by_key[r["key"]] = r
    cols = (result.get("table") or {}).get("columns") or TABLE_COLUMNS
    result["table"] = {"columns": cols, "rows": [[r.get(c) for c in cols] for r in result.get("sites") or []]}
    for f in (result.get("geojson") or {}).get("features") or []:
        r = by_key.get(f["properties"].get("key"))
        if r:
            f["properties"].update(area_km2=r.get("area_km2"), q100_per_km2=r.get("q100_per_km2"))
    return result


# ── downloads ───────────────────────────────────────────────────────────────


def to_csv(result: dict[str, Any]) -> str:
    """The per-site table as CSV."""
    table = result.get("table") or {}
    return pd.DataFrame(table.get("rows") or [], columns=table.get("columns") or TABLE_COLUMNS).to_csv(index=False)


def to_xlsx(result: dict[str, Any]) -> bytes:
    """A workbook: the per-site table, the regional growth curve, field significance and notes. Needs openpyxl."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sites"
    table = result.get("table") or {}
    ws.append(list(table.get("columns") or TABLE_COLUMNS))
    for r in table.get("rows") or []:
        ws.append([v if not isinstance(v, (list, dict)) else str(v) for v in r])
    ws.freeze_panes = "A2"

    rfa = result.get("regional_frequency") or {}
    g = wb.create_sheet("Regional growth curve")
    if rfa.get("regional"):
        reg = rfa["regional"]
        g.append(["Return period (years)", "Growth factor (x index flood)"])
        for t, f in (reg.get("growth_curve") or {}).items():
            g.append([float(t), f])
        g.append([])
        g.append(["L-CV", reg.get("l_cv")])
        g.append(["L-skewness", reg.get("l_skew")])
        g.append(["L-kurtosis", reg.get("l_kurt")])
        het = rfa.get("heterogeneity") or {}
        g.append(["Heterogeneity H", het.get("H")])
        g.append(["Class", het.get("class")])
        g.append(["Sites in pool", rfa.get("n_sites")])
    else:
        g.append([rfa.get("error") or "No regional growth curve."])

    fs = result.get("field_significance") or {}
    f = wb.create_sheet("Field significance")
    counts = fs.get("counts") or {}
    for label, value in (("Sites tested", fs.get("n_tested")), ("Up", counts.get("up")),
                         ("Down", counts.get("down")), ("None", counts.get("none")),
                         ("Expected by chance", fs.get("expected_by_chance")),
                         ("Significant after FDR", (fs.get("fdr") or {}).get("n_significant")),
                         ("Walker p_min", (fs.get("walker") or {}).get("p_min")),
                         ("Walker critical", (fs.get("walker") or {}).get("critical")),
                         ("Binomial p", fs.get("binomial_p")), ("Verdict", fs.get("verdict")),
                         ("Caveat", fs.get("caveat"))):
        f.append([label, value])

    n = wb.create_sheet("Notes")
    n.append(["Question", result.get("question")])
    n.append(["Headline", result.get("headline")])
    for note in result.get("notes") or []:
        n.append(["Note", note])
    for m in result.get("methods") or []:
        n.append(["Method", m.get("name"), m.get("citation")])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = [
    "MAX_LIVE_FETCHES", "MAX_SITES", "METHODS", "TABLE_COLUMNS", "analyse_sites", "apply_areas", "benjamini_hochberg",
    "field_significance", "gather_series", "index_flood_rfa", "inventory", "site_summary", "study_area",
    "to_csv", "to_xlsx",
]
