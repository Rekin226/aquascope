"""The Excel workbook of a study: README, the inventory, the plan, the gates, one sheet per table, the figure index,
the token ledger.

openpyxl is imported inside :func:`workbook_bytes`, never at module import.
"""

from __future__ import annotations

import io
import json
import re
from typing import Any

from aquascope.studio.deliverables import _common as c
from aquascope.studio.deliverables.tables import frame_of
from aquascope.studio.workspace import Workspace

_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_MAX_ROWS = 100_000


def sheet_title(name: str, taken: set[str]) -> str:
    """A sheet title openpyxl accepts: no forbidden characters, at most 31 characters, unique among ``taken``."""
    base = _BAD_SHEET_CHARS.sub("_", name).strip() or "sheet"
    base = base[:31]
    title, n = base, 2
    while title.lower() in {t.lower() for t in taken}:
        suffix = f"_{n}"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    taken.add(title)
    return title


def _cell(v: Any) -> Any:
    """A value openpyxl can store: scalars as they are, everything else as JSON text."""
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    try:
        return json.dumps(v, ensure_ascii=False, default=str)
    except (TypeError, ValueError):
        return str(v)


def _fill(sheet: Any, columns: list[str], rows: list[list[Any]], *, freeze: bool = True) -> None:
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    sheet.append([str(col) for col in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for r in rows[:_MAX_ROWS]:
        sheet.append([_cell(v) for v in r])
    if freeze:
        sheet.freeze_panes = "A2"
    widths = [len(str(col)) for col in columns]
    for r in rows[:200]:
        for i, v in enumerate(r[: len(widths)]):
            widths[i] = max(widths[i], min(len(str(v)) if v is not None else 0, 60))
    for i, w in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = max(8, min(w + 2, 62))


def _readme_rows(ws: Workspace) -> list[list[Any]]:
    rows: list[list[Any]] = [
        ["Title", c.title_of(ws)],
        ["Answer", c.answer_of(ws)],
        ["Site", c.site_text(ws)],
        ["Problem", ws.brief.problem],
        ["Decision", ws.brief.decision or ""],
        ["Date", c.date_of(ws)],
        ["Status", ws.status],
        ["AquaScope version", c.version()],
        ["Model", c.model_line(ws)],
        ["Workspace id", ws.id],
        ["", ""],
        ["Key numbers", ""],
    ]
    for k in c.key_numbers(ws):
        unit = f" {k.get('unit')}" if k.get("unit") else ""
        step = f" (step {k.get('step')})" if k.get("step") else ""
        rows.append([str(k.get("label") or ""), f"{k.get('value')}{unit}{step}"])
    rows += [
        ["", ""],
        ["Sheets", "Inventory (the datasets), Plan (the steps), Gates (every gate outcome), one sheet per result "
                   "table, Figures (the figure index), Ledger (tokens per role)"],
        ["Reproduce", "aquascope run study.yaml (the study.yaml in the bundle), or open study.ipynb"],
        ["How to cite", c.citation()],
    ]
    return rows


def _inventory_rows(ws: Workspace) -> tuple[list[str], list[list[Any]]]:
    cols = ["id", "kind", "variable", "source", "station_id", "name", "lat", "lon", "distance_km", "start", "end",
            "years", "resolution", "n", "note", "quality"]
    rows = []
    for d in (ws.inventory.datasets if ws.inventory else []):
        dd = d.to_dict()
        rows.append([dd.get(k) for k in cols])
    return cols, rows


def _plan_rows(ws: Workspace) -> tuple[list[str], list[list[Any]]]:
    cols = ["id", "tool", "arguments", "method", "gates", "rationale", "depends_on", "fallback", "outputs"]
    rows = []
    for i, s in enumerate(c.steps_of(ws), 1):
        rows.append([s.get("id") or f"s{i}", s.get("tool"), s.get("arguments") or {}, s.get("method"),
                     s.get("expects") or [], s.get("rationale"), s.get("depends_on") or [], s.get("fallback"),
                     s.get("outputs") or []])
    return cols, rows


def _gate_rows(ws: Workspace) -> tuple[list[str], list[list[Any]]]:
    cols = ["step", "check", "passed", "detail", "path", "value"]
    rows = [[g.get("step"), g.get("check"), g.get("passed"), g.get("detail"), g.get("path") or g.get("paths"),
             g.get("value")] for g in c.gates_of(ws)]
    return cols, rows


def workbook_bytes(ws: Workspace) -> bytes:
    """The workbook as ``.xlsx`` bytes.

    Sheets: ``README`` (title, answer, site, date, version, key numbers, how to cite), ``Inventory``, ``Plan``,
    ``Gates``, one sheet per table artifact (titles at most 31 characters, unique), ``Figures`` (id, file,
    caption, step), ``Ledger`` (tokens per role). Header rows are bold and frozen; column widths follow the text.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    taken: set[str] = set()

    readme = wb.active
    readme.title = sheet_title("README", taken)
    readme["A1"] = c.title_of(ws)
    readme["A1"].font = Font(bold=True, size=14)
    for row in _readme_rows(ws):
        readme.append(row)
    for cell in readme["A"]:
        cell.font = Font(bold=True)
    for cell in readme["B"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 20
    readme.column_dimensions["B"].width = 100

    cols, rows = _inventory_rows(ws)
    _fill(wb.create_sheet(sheet_title("Inventory", taken)), cols, rows)
    cols, rows = _plan_rows(ws)
    _fill(wb.create_sheet(sheet_title("Plan", taken)), cols, rows)
    cols, rows = _gate_rows(ws)
    _fill(wb.create_sheet(sheet_title("Gates", taken)), cols, rows)
    findings = ws.findings or {}
    if findings.get("findings") or findings.get("decision"):
        frows = [[f.get("id"), f.get("grade"), f.get("claim"), "; ".join(f.get("basis") or []),
                  f.get("evidence") or {}]
                 for f in findings.get("findings") or []]
        d = findings.get("decision") or {}
        if d:
            band = d.get("band")
            frows.append(["decision", d.get("grade"), d.get("answer"),
                          f"value {d.get('value')} {d.get('unit') or ''}; band {band}".strip(),
                          d.get("evidence") or {}])
        _fill(wb.create_sheet(sheet_title("Findings", taken)), ["id", "grade", "claim", "basis", "evidence"], frows)

    for a in ws.artifacts_of("table"):
        stem = a.name.rsplit("/", 1)[-1].removesuffix(".csv") or a.id
        sheet = wb.create_sheet(sheet_title(stem, taken))
        try:
            df = frame_of(a)
        except Exception:  # noqa: BLE001 - an unreadable CSV gets a note, not a crash
            sheet.append([f"{a.id}: the CSV could not be read"])
            continue
        columns, values = c.frame_cells(df)
        _fill(sheet, columns, values)
        if a.caption:
            sheet.cell(row=1, column=len(df.columns) + 2, value=a.caption).font = Font(italic=True)

    figs = [[a.id, a.name, a.caption, a.step, (a.meta or {}).get("kind"), a.media_type, a.size]
            for a in ws.figures()]
    _fill(wb.create_sheet(sheet_title("Figures", taken)), ["id", "file", "caption", "step", "kind", "media_type",
                                                            "bytes"], figs)

    ledger = [[role, v.get("calls", 0), v.get("prompt_tokens", 0), v.get("completion_tokens", 0),
               v.get("prompt_tokens", 0) + v.get("completion_tokens", 0)] for role, v in ws.ledger.items()]
    ledger.append(["total", sum(r[1] for r in ledger), sum(r[2] for r in ledger), sum(r[3] for r in ledger),
                   sum(r[4] for r in ledger)])
    _fill(wb.create_sheet(sheet_title("Ledger", taken)), ["role", "calls", "prompt_tokens", "completion_tokens",
                                                          "tokens"], ledger)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
