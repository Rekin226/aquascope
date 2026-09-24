"""Study this area: the multi-gauge engine on synthetic series, with the Archive and the agencies mocked."""

from __future__ import annotations

import json

import numpy as np
import pandas as pd
import pytest

from aquascope import area_study as aa


def _daily(seed: int, years: int = 30, scale: float = 10.0, trend: float = 0.0, start: int = 1990) -> pd.Series:
    """A daily flow series whose annual peaks follow a Gumbel, with an optional linear trend in the peaks."""
    rng = np.random.default_rng(seed)
    idx = pd.date_range(f"{start}-01-01", f"{start + years - 1}-12-31", freq="D")
    base = pd.Series(scale * 0.2 * rng.random(len(idx)), index=idx)
    for i, y in enumerate(range(start, start + years)):
        peak = scale * (1 + 0.3 * rng.gumbel()) + trend * i
        base[pd.Timestamp(f"{y}-06-15")] = max(peak, scale * 0.5)
    return base


def _station(i: int, source: str = "usgs", **kw) -> dict:
    return {"source": source, "station_id": f"S{i:02d}", "name": f"Gauge {i}", "lat": 40 + i * 0.01,
            "lon": -75 - i * 0.01, "variables": ["discharge"], "period_start": f"{1990 - i}-01-01", **kw}


def test_benjamini_hochberg_matches_the_textbook_step_up():
    p = [0.01, 0.04, 0.03, 0.005, 0.5]
    # sorted: .005 (<=.01), .01 (<=.02), .03 (<=.03), .04 (> .04? no: .04 <= .04) -> the four smallest pass
    assert aa.benjamini_hochberg(p, q=0.05) == [True, True, True, True, False]
    assert aa.benjamini_hochberg([0.2, 0.3, 0.9]) == [False, False, False]
    assert aa.benjamini_hochberg([]) == []


def test_field_significance_counts_and_walker():
    sites = {f"s{i}": {"trend_p": 0.5, "trend": "none"} for i in range(8)}
    sites["s0"] = {"trend_p": 0.0001, "trend": "up"}
    sites["s1"] = {"trend_p": 0.001, "trend": "up"}
    sites["s2"] = {"trend_p": 0.03, "trend": "down"}
    sites["u"] = {"trend_p": None, "trend": "untested"}
    fs = aa.field_significance(sites)
    assert fs["n_tested"] == 8 and fs["n_untested"] == 1
    assert fs["counts"] == {"up": 2, "down": 1, "none": 5}
    assert set(fs["fdr"]["significant"]) == {"s0", "s1"}
    assert fs["walker"]["significant"] and fs["field_significant"]
    assert "upward" in fs["verdict"] and "independent" in fs["caveat"]


def test_field_significance_with_nothing_testable_says_so():
    fs = aa.field_significance({"a": {"trend_p": None, "trend": "untested"}})
    assert fs["n_tested"] == 0 and fs["field_significant"] is None and "years" in fs["verdict"]


def test_index_flood_rfa_on_a_homogeneous_region():
    rng = np.random.default_rng(3)
    amax = {f"s{i}": (50 + 20 * i) * (1 + 0.25 * rng.gumbel(size=40)) for i in range(8)}
    amax["short"] = np.array([1.0, 2.0, 3.0])
    rfa = aa.index_flood_rfa(amax, n_sim=200, seed=1)
    assert rfa["n_sites"] == 8 and rfa["left_out"] == ["short"]
    g = rfa["regional"]["growth_curve"]
    assert g["2"] < 1.0 < g["10"] < g["100"]  # the median flood is below the mean, the 100-year above it
    assert rfa["heterogeneity"]["H"] is not None and rfa["heterogeneity"]["H"] < 2
    s0 = rfa["sites"]["s0"]
    assert s0["q_regional"]["100"] == pytest.approx(s0["index_flood"] * g["100"], rel=1e-3)
    assert "discordancy" in s0 and rfa["discordancy"]["critical"] == pytest.approx(2.140)


def test_index_flood_rfa_flags_heterogeneity_when_l_cv_differs():
    rng = np.random.default_rng(5)
    amax = {f"low{i}": 100 * (1 + 0.05 * rng.gumbel(size=40)) for i in range(4)}
    amax |= {f"high{i}": 100 * np.exp(0.9 * rng.normal(size=40)) for i in range(4)}
    rfa = aa.index_flood_rfa(amax, n_sim=200, seed=1)
    assert rfa["heterogeneity"]["H"] >= 2 and rfa["heterogeneity"]["class"] == "definitely heterogeneous"
    assert any("growth curve may not fit" in n for n in rfa["notes"])


def test_index_flood_rfa_needs_two_long_records():
    rfa = aa.index_flood_rfa({"a": np.arange(1, 30), "b": np.arange(1, 5)})
    assert "error" in rfa and rfa["n_sites"] == 1


def test_inventory_from_a_bbox_keeps_one_per_site_and_the_variable():
    rows = [
        {"source": "usgs", "station_id": "A", "latitude": 40.0, "longitude": -75.0, "variables": ["discharge"],
         "period_start": "1950-01-01"},
        {"source": "usgs", "station_id": "A2", "site_id": "A", "latitude": 40.0, "longitude": -75.0,
         "variables": ["discharge"], "period_start": "2000-01-01"},
        {"source": "usgs", "station_id": "B", "latitude": 40.5, "longitude": -75.5, "variables": ["water_level"]},
        {"source": "usgs", "station_id": "C", "latitude": 50.0, "longitude": -75.0, "variables": ["discharge"]},
    ]
    inv = aa.inventory(bbox=(-76, 39, -74, 41), rows=rows)
    assert [s["station_id"] for s in inv["sites"]] == ["A"]
    assert inv["n_in_area"] == 3 and inv["n_without_variable"] == 1
    assert any("duplicate" in n for n in inv["notes"])


def test_inventory_caps_the_sites_longest_first():
    inv = aa.inventory([_station(i) for i in range(5)], max_sites=2)
    assert [s["station_id"] for s in inv["sites"]] == ["S04", "S03"]
    assert inv["n_over_cap"] == 3 and len(inv["dropped"]) == 3


def test_study_area_reads_the_archive_first_and_caps_live_fetches():
    stations = [_station(i) for i in range(6)]
    in_archive = {"S00", "S01", "S02"}
    calls = {"archive": [], "live": []}

    def archive(source, sid, variable):
        calls["archive"].append(sid)
        return _daily(int(sid[1:]), trend=0.0) if sid in in_archive else None

    def live(source, sid, variable, period_start=None):
        calls["live"].append(sid)
        return _daily(int(sid[1:]) + 10)

    events = []
    res = aa.study_area(stations, archive_reader=archive, live_reader=live, max_live=2, n_sim=100,
                       on_progress=events.append)
    assert len(calls["archive"]) == 6 and len(calls["live"]) == 2
    s = res["summary"]
    assert (s["n_archive"], s["n_live"], s["n_skipped"], s["n_studied"]) == (3, 2, 1, 5)
    skipped = [r for r in res["sites"] if r["status"] == "skipped"]
    assert len(skipped) == 1 and "live-fetch cap" in skipped[0]["note"]
    assert any("skipped" in n for n in res["notes"])
    assert {e["phase"] for e in events} >= {"inventory", "fetch", "analyse", "regional"}
    # everything the page reads is JSON
    json.dumps(res, allow_nan=False)
    assert res["geojson"]["features"][0]["geometry"]["type"] == "Point"
    props = res["geojson"]["features"][0]["properties"]
    assert {"trend", "q100", "q100_per_km2", "status"} <= set(props)
    assert res["table"]["columns"] == aa.TABLE_COLUMNS
    assert res["regional_frequency"]["n_sites"] == 5
    assert res["headline"].startswith("5 of 6 gauges studied")


def test_study_area_reports_a_failed_live_fetch_without_sinking():
    def archive(*_a):
        return None

    def live(source, sid, variable, period_start=None):
        if sid == "S01":
            raise RuntimeError("HTTP 429 Too Many Requests")
        return _daily(int(sid[1:]), trend=2.0)

    res = aa.study_area([_station(i) for i in range(3)], archive_reader=archive, live_reader=live, n_sim=50)
    failed = [r for r in res["sites"] if r["status"] == "failed"]
    assert len(failed) == 1 and "429" in failed[0]["note"]
    assert res["summary"]["n_studied"] == 2


def test_upward_trends_show_in_the_field():
    stations = [_station(i) for i in range(6)]

    def archive(source, sid, variable):
        return _daily(int(sid[1:]), trend=1.5)

    res = aa.study_area(stations, archive_reader=archive, n_sim=50)
    fs = res["field_significance"]
    assert fs["counts"]["up"] >= 4 and fs["field_significant"]
    ups = [r for r in res["sites"] if r["trend"] == "up"]
    assert all(r["sens_slope_per_year"] > 0 for r in ups)


def test_q100_per_km2_uses_the_area_when_known():
    s = _daily(1)
    out = aa.site_summary(s, area_km2=250.0)
    assert out["q100"] and out["q100_per_km2"] == pytest.approx(out["q100"] / 250.0, rel=1e-3)
    assert aa.site_summary(s)["q100_per_km2"] is None


def test_short_record_gets_no_trend_and_a_note():
    stations = [_station(0)]
    res = aa.study_area(stations, archive_reader=lambda *a: _daily(0, years=5), n_sim=10)
    row = res["sites"][0]
    assert row["trend"] == "untested" and "too short" in row["note"]
    assert "error" in res["regional_frequency"]


def test_downloads():
    res = aa.study_area([_station(i) for i in range(3)], archive_reader=lambda s, sid, v: _daily(int(sid[1:])),
                       n_sim=20)
    csv = aa.to_csv(res)
    assert csv.splitlines()[0].startswith("source,station_id,name") and len(csv.splitlines()) == 4
    pytest.importorskip("openpyxl")
    data = aa.to_xlsx(res)
    assert data[:2] == b"PK"
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Sites", "Regional growth curve", "Field significance", "Notes"]


def test_apply_areas_fills_q100_per_km2_everywhere():
    res = aa.study_area([_station(i) for i in range(3)], archive_reader=lambda s, sid, v: _daily(int(sid[1:])),
                        n_sim=20)
    assert all(r["q100_per_km2"] is None for r in res["sites"])
    aa.apply_areas(res, {"usgs/S00": 100.0, "usgs/S01": None, "usgs/S02": "bad"})
    row = next(r for r in res["sites"] if r["key"] == "usgs/S00")
    assert row["q100_per_km2"] == pytest.approx(row["q100"] / 100.0, rel=1e-3)
    col = res["table"]["columns"].index("q100_per_km2")
    assert [r[col] for r in res["table"]["rows"]].count(None) == 2
    feat = next(f for f in res["geojson"]["features"] if f["properties"]["key"] == "usgs/S00")
    assert feat["properties"]["area_km2"] == 100.0
