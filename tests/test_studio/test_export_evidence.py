"""Exports must retain analyzed inputs and distinguish every fitted uncertainty."""
import io

import pandas as pd

from aquascope.studio.deliverables import report_docx, tables
from aquascope.studio.deliverables._payload import return_levels_of
from aquascope.studio.roles import analysts


def test_long_subdaily_input_survives_study_table(monkeypatch):
    import aquascope.explore

    observed = pd.Series([0.123456789012345 + i for i in range(26002)],
                         index=pd.date_range("2000-01-01T00:15:00", periods=26002, freq="h", tz="UTC"))
    observed.iloc[5] = float("nan")

    def analyze(*args, store, **kwargs):
        store["series"] = observed
        return {"n": len(observed.dropna()), "series_downsampled": True,
                "series": {"t": ["2000-01-01"], "v": [12.0]}}

    monkeypatch.setattr(aquascope.explore, "analyze_station", analyze)
    result = analysts.analyze_station_full("usgs", "USGS-01013500")
    artifact = tables.tables_for("s2", "analyze_station", result, names=["series"])[0]
    actual = pd.read_csv(io.BytesIO(artifact.data), float_precision="round_trip")
    assert len(actual) == result["n"] == 26001
    assert actual.datetime.tolist() == [t.isoformat() for t in observed.dropna().index]
    assert actual.value.tolist() == observed.dropna().tolist()
    assert result["series"]["v"] == [12.0], "the figure data remains independent"


def test_unavailable_full_record_is_not_silently_replaced_by_plot():
    payload = {"series_downsampled": True, "series": {"t": ["2020-01-01"], "v": [5.0]}}
    assert tables.tables_for("s1", "analyze_station", payload, names=["series"]) == []


def test_flood_step_retains_its_own_observations(monkeypatch):
    observation = {"t": ["2001-01-01T00:00:00"], "v": [7.123456789]}
    payload = {"source": "usgs", "station_id": "gauge", "data_snapshot": "sha256:own-input",
               "variable": "discharge", "unit": "m3/s", "n": 1, "observations": observation,
               "ffa": {"return_periods": [100], "fits": {"gev_lmoments": {"q": [17.]}}}}
    monkeypatch.setattr(analysts, "analyze_station_full", lambda *a, **k: payload)
    result = analysts.flood_frequency_full("usgs", "gauge")
    assert result["observations"] == observation and result["data_snapshot"] == "sha256:own-input"
    artifacts = tables.tables_for("s3", "flood_frequency", result)
    retained = next(a for a in artifacts if a.id == "tab-s3-series")
    assert b"2001-01-01T00:00:00,7.123456789" in retained.data


def test_uploaded_subdaily_timestamps_are_not_truncated():
    payload = {"series": {"t": ["2020-01-01T01:15:00+00:00", "2020-01-01T02:15:00+00:00"],
                          "v": [1.1234567890123, 2.0]}}
    columns, rows = tables.make("series", payload)
    assert columns == ["datetime", "value"]
    assert [r[0] for r in rows] == payload["series"]["t"]
    assert rows[0][1] == payload["series"]["v"][0]


def test_each_exported_interval_stays_with_its_fit():
    payload = {"unit": "m3/s", "ffa": {"return_periods": [100], "fits": {
        "gev_lmoments": {"q": [583.2]},
        "lp3": {"q": [570.0], "ci": [[510.0, 630.0]], "ci_level": .9,
                "interval_method": "variance_of_estimate"},
        "gev_bootstrap": {"q": [539.2], "ci": [[403.8, 767.1]], "ci_level": .95,
                          "estimator": "gev_mle", "interval_method": "percentile_bootstrap"}}}}
    artifact = tables.tables_for("s3", "flood_frequency", payload, names=["return_levels"])[0]
    records = tables.frame_of(artifact).set_index("estimator")
    assert records.loc["gev_lmoments", "estimate"] == 583.2
    assert pd.isna(records.loc["gev_lmoments", "lower"])
    assert records.loc["lp3", "lower"] == 510.0
    assert records.loc["gev_mle", "estimate"] == 539.2
    assert records.loc["gev_mle", "lower"] == 403.8
    assert records.loc["gev_mle", "confidence_level"] == .95
    assert "95 %" in return_levels_of(payload)["band"]
    del payload["ffa"]["fits"]["gev_bootstrap"]["ci_level"]
    assert "unrecorded" in return_levels_of(payload)["band"]


def test_word_tables_and_code_are_native_readable_blocks():
    from docx import Document

    doc = Document()
    report_docx.markdown_to_docx(doc, "| Station | Unit |\n| --- | --- |\n| Fish River | m3/s |\n\n"
                               "```json\n{\n  \"result\": 463.9\n}\n```")
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[1].cells[0].text == "Fish River"
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "| ---" not in text and "```" not in text
    assert '{\n  "result": 463.9\n}' in text


def test_wide_word_tables_repeat_identifying_columns():
    from docx import Document

    doc = Document()
    columns = ["T", "estimator", "unit", "estimate", "lower", "upper", "method", "level"]
    report_docx._add_table(doc, columns, [[100, "gev_mle", "m3/s", 539.2, 403.8, 767.1, "bootstrap", .95]], None)
    assert len(doc.tables) == 2
    for table in doc.tables:
        assert [c.text for c in table.rows[0].cells[:2]] == ["T", "estimator"]
        assert [c.text for c in table.rows[1].cells[:2]] == ["100", "gev_mle"]


def test_workbook_preserves_observations_above_the_old_row_cap():
    from openpyxl import load_workbook

    from aquascope.studio.deliverables.workbook import workbook_bytes
    from aquascope.studio.workspace import Artifact, Workspace

    count = 100001
    ws = Workspace()
    ws.artifacts = [Artifact(id="tab-s1-series", kind="table", name="tables/s1_series.csv",
                             media_type="text/csv", data=("row,value\n" + "".join(f"{i},{i}.25\n"
                             for i in range(count))).encode(), meta={"name": "series", "rows": count})]
    book = load_workbook(io.BytesIO(workbook_bytes(ws)), read_only=True)
    sheets = [s for s in book if s.title.startswith("s1_series")]
    assert len(sheets) == 2
    rows = [row for sheet in sheets for row in sheet.iter_rows(min_row=2, max_col=2, values_only=True)]
    assert len(rows) == count
    assert rows[0] == (0, .25) and rows[-1] == (100000, 100000.25)
    assert [r[0] for r in rows] == list(range(count))


def test_workbook_does_not_mark_a_skipped_gate_as_passed():
    from aquascope.studio.deliverables.workbook import _gate_rows
    from aquascope.studio.workspace import Workspace

    ws = Workspace()
    ws.run = {'gates': [{'step': 's4', 'check': 'cross_check_ratio', 'passed': True, 'skipped': True}]}
    columns, rows = _gate_rows(ws)
    row = dict(zip(columns, rows[0]))
    assert row['skipped'] is True and row['passed'] is None
