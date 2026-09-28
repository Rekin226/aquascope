"""Execute exported notebook cells: new results must never inherit old evidence."""
import functools
import io

import pandas as pd
from openpyxl import load_workbook

from aquascope.studio.deliverables import notebook
from aquascope.studio.workspace import Artifact, Workspace
from aquascope.study import Step, Study


def original_files(tmp_path):
    ws = Workspace(site={'lat': 47.2, 'lon': -68.5})
    ws.brief.problem = 'Inspect the available flow record'
    ws.tables = {'upload:retained': 'date,value\n2020-01-01,5\n'}
    study = Study(question=ws.brief.problem, steps=[Step(
        id='s1', tool='analyze_station', arguments={'source': 'usgs', 'station_id': 'A'},
        expects=[{'check': 'not_empty', 'path': 'stats'}])])
    ws.set_study(study)
    ws.report = {'answer': 'STALE ANSWER', 'key_numbers': [{'label': 'Old mean', 'value': 999999}]}
    ws.findings = {'decision': {'answer': 'STALE DECISION'}}
    ws.artifacts = [Artifact('old-table', 'table', 'tables/old.csv', b'value\n999999\n')]
    ws.ledger = {'author': {'calls': 3, 'prompt_tokens': 900}}
    (tmp_path / 'study.yaml').write_text(study.to_yaml())
    (tmp_path / 'workspace.json').write_text(ws.to_json())
    (tmp_path / 'workbook.xlsx').write_bytes(b'original workbook stays untouched')
    return ws


def test_exported_notebook_executes_fresh_evidence_without_overwriting_original(tmp_path, monkeypatch):
    original = original_files(tmp_path)
    before = (tmp_path / 'workspace.json').read_bytes()
    dates = pd.date_range('2020-01-01', periods=501)
    payload = {'source': 'usgs', 'station_id': 'A', 'variable': 'discharge', 'unit': 'm3/s',
               'n': 501, 'years': 1.4, 'start': '2020-01-01', 'end': dates[-1].date().isoformat(),
               'stats': {'mean': 9.125}, 'series_downsampled': True,
               'series': {'t': ['2020-01-01'], 'v': [9.]},
               'observations': {'t': [d.isoformat() for d in dates], 'v': [9.125] * 501}}
    rerun = notebook.rerun_workspace
    monkeypatch.setattr(notebook, 'rerun_workspace', functools.partial(
        rerun, tools={'analyze_station': lambda **kw: payload}))
    monkeypatch.chdir(tmp_path)
    namespace = {}
    for cell in notebook.notebook_cells(original):
        if cell['cell_type'] == 'code':
            exec(compile(cell['source'], '<exported-notebook>', 'exec'), namespace)
    ws = namespace['ws']
    assert ws.id != original.id and ws.tables == original.tables
    assert not ws.ledger and ws.model is None
    assert ws.run['gates'] and ws.run['gates'][0]['passed'] is True
    assert 'STALE' not in str(ws.report) and '999999' not in str(ws.report)
    assert ws.artifact('old-table') is None
    table = pd.read_csv(io.BytesIO(ws.artifact('tab-s1-series').data))
    assert len(table) == 501 and table.value.tolist() == [9.125] * 501
    book = load_workbook(namespace['out'] / 'workbook.xlsx', read_only=True)
    assert book['Gates']['B2'].value == 'not_empty'
    assert book['Gates']['C2'].value is True
    assert all('Old mean' not in str(row) for row in book['README'].values)
    assert (namespace['out'] / 'workspace.json').exists()
    assert (tmp_path / 'workspace.json').read_bytes() == before
    assert (tmp_path / 'workbook.xlsx').read_bytes() == b'original workbook stays untouched'


def test_failed_replay_does_not_reuse_original_claims_or_artifacts(tmp_path):
    original_files(tmp_path)

    def unavailable(**kwargs):
        raise RuntimeError('agency unavailable')

    ws = notebook.rerun_workspace(tmp_path / 'study.yaml', tmp_path / 'workspace.json',
                                  tools={'analyze_station': unavailable})
    assert ws.run['results'][0]['ok'] is False
    assert ws.artifact('old-table') is None and ws.artifact('tab-s1-series') is None
    assert 'STALE' not in str(ws.report)
    assert ws.report['not_established']


def test_notebook_install_pins_recorded_source(tmp_path, monkeypatch):
    ws = original_files(tmp_path)
    monkeypatch.setenv('AQUASCOPE_REVISION', '104ab2d')
    intro = notebook.notebook_cells(ws)[0]['source']
    assert 'aquascope.git@104ab2d' in intro and 'current data' in intro
