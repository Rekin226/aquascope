"""The accuracy plumbing (#414) and the documents' figures and tables (#415)."""

from __future__ import annotations

import io

from aquascope.studio.model import Model
from aquascope.studio.roles import author, critic
from aquascope.studio.workspace import Artifact, Workspace
from tests.test_studio.conftest import FakeModel
from tests.test_studio.test_critic_author import _ran
from tests.test_studio.test_methodologist import VALID_PLAN

BRIEF = {"decision": "size the crossing", "quantities": ["the 100-year flow with a band"], "kind": "flood_risk",
         "playbook": "flood_risk", "intake": {"return_period": 100}, "assumptions": [], "questions": [],
         "ready": True}
GOOD = "About 520 m3/s at uk_ea 3400TH by GEV L-moments."


def test_the_keyed_authors_prose_is_sentence_checked_and_the_drop_is_counted(studio_factory):
    client = FakeModel({
        "consultant": [BRIEF], "methodologist": [VALID_PLAN],
        "author": [{"title": "Design flow", "answer": GOOD + " The catchment drains 77777 km2.",
                    "sections": {"summary": "The flow is 520 m3/s at uk_ea 3400TH. It rained 31415 mm in 1066."}}],
        "critic": [{"issues": []}],
    })
    s, _ = studio_factory(client=client)
    s.say("A culvert on the Thames at Kingston, 100-year")
    r = s.approve()
    ws = s.workspace
    assert r.kind == "report" and "About 520" in ws.report["answer"] and "77777" not in ws.report["answer"]
    summary = next(x["text"] for x in ws.report["sections"] if x["id"] == "summary")
    assert "31415" not in summary and "520" in summary
    assert ws.report["dropped"] == 2 and r.payload["dropped"] == 2
    assert "2 sentence(s) were dropped by the checks" in author.footer_line(ws)
    assert any(e["event"] == "dropped" and e["role"] == "critic" for e in ws.events)


def test_a_fix_earns_one_rewrite_a_modelled_second_critique_and_a_notice_when_still_not_ok(studio_factory):
    client = FakeModel({
        "consultant": [BRIEF], "methodologist": [VALID_PLAN],
        "author": [{"title": "t", "answer": GOOD, "sections": {"summary": GOOD}},
                   {"title": "t", "answer": GOOD, "sections": {"summary": GOOD + " Rewritten."}}],
        "critic": [{"issues": [{"section": "summary", "severity": "fix", "text": "x", "fix": "y"}]},
                   {"issues": [{"section": "summary", "severity": "fix", "text": "still", "fix": "z"}]}],
    })
    s, _ = studio_factory(client=client)
    s.say("A culvert on the Thames at Kingston, 100-year")
    r = s.approve()
    ws = s.workspace
    assert len(client.calls("author")) == 2 and len(client.calls("critic")) == 2, "fix round, modelled critique"
    assert any(e["event"] == "fixes" for e in ws.events)
    assert ws.report["answer"].startswith("Notice:") and "summary" in ws.report["notice"]
    assert r.payload["critique_ok"] is False and ws.report["critique"]["ok"] is False
    assert "Rewritten." in next(x["text"] for x in ws.report["sections"] if x["id"] == "summary")
    assert "Notice:" in author.to_markdown(ws).splitlines()[0:12].__str__()


def test_every_failed_check_is_a_fix_for_the_author_with_the_repair_named():
    critique = {"ok": False, "issues": [{"section": "summary", "severity": "note", "text": "n", "fix": ""}],
                "checks": [{"name": "numbers_come_from_tools", "passed": False, "detail": "77777 is in no result"},
                           {"name": "trend_matches_the_test", "passed": False, "detail": "p = 0.302 called sig."},
                           {"name": "record_is_named", "passed": True, "detail": ""}]}
    fixes = critic.fixes_for(critique)
    assert [f["check"] for f in fixes] == ["numbers_come_from_tools", "trend_matches_the_test"]
    assert all(f["severity"] == "fix" for f in fixes) and "p is below 0.05" in fixes[1]["fix"]
    assert critic.failed_checks(critique) == ["numbers_come_from_tools", "trend_matches_the_test"]
    line = critic.notice(critique)
    assert line.startswith("Notice:") and "trend_matches_the_test" in line and not any(ch.isdigit() for ch in line)
    assert critic.notice({"ok": True, "checks": [], "issues": []}) is None


def test_a_keyless_report_with_a_failed_check_opens_with_the_notice(no_deliverables):
    ws = _ran()
    author.author_report(ws, None)
    ws.report["answer"] = "The 100-year flow is 77777 m3/s at Kingston (uk_ea 3400TH), with a 90 % band."
    ws.report["sections"][0]["text"] = ws.report["answer"]
    out = critic.critique(ws, None)
    assert not out["ok"] and "numbers_come_from_tools" in out["failed"]
    fixes = critic.fixes_for(out)
    author.author_report(ws, None, issues=fixes)
    assert "77777" not in ws.report["answer"], "the template repair drops the sentence the check refuses"


def test_small_p_values_print_as_a_threshold():
    assert author.p_text(4.3e-07) == "< 0.001" and author.p_text(0.302) == 0.302 and author.p_text(None) is None
    assert author._small_p("no trend (p = 4.3e-07, 40 years) and p = 0.302 elsewhere") == \
        "no trend (p < 0.001, 40 years) and p = 0.302 elsewhere"
    from aquascope.studio.model import compact

    assert compact({"p_value": 4.3e-07})["p_value"] == 4.3e-07 and compact({"x": 1.23456789})["x"] == 1.234568


def test_the_ledger_prices_a_known_model_and_the_ceiling_switches_the_roles_to_keyless():
    ws = Workspace()
    ws.site = {"lat": 51.4, "lon": -0.3}
    client = FakeModel({"author": ["one", "two", "three"]})
    model = Model.resolve(ws, client=client, model="claude-haiku-4-5", provider="custom", max_usd=0.0004)
    assert model.call("author", "You are the Author", {}) == "one"
    per_call = 120 * 1.0 / 1e6 + 30 * 5.0 / 1e6
    assert ws.ledger["author"]["cost_usd"] == round(per_call, 6) and ws.total_usd == round(per_call, 6)
    assert model.call("author", "You are the Author", {}) == "two"
    assert ws.budget is not None and ws.budget["max_usd"] == 0.0004 and ws.budget["role"] == "author"
    assert model.call("author", "You are the Author", {}) is None, "past the ceiling, keyless behaviour"
    assert any(e["event"] == "model_skipped" for e in ws.events) and any(e["event"] == "budget" for e in ws.events)
    assert ws.ledger["author"]["calls"] == 2 and ws.total_usd == round(2 * per_call, 6)
    assert "The spend ceiling of" in author.footer_line(ws)
    unknown = Workspace()
    unknown.site = {"lat": 0, "lon": 0}
    m2 = Model.resolve(unknown, client=FakeModel({"author": ["x"]}), model="fake", provider="custom", max_usd=0.001)
    assert m2.call("author", "You are the Author", {}) == "x" and unknown.total_usd is None
    assert any("cannot be enforced" in e["detail"] for e in unknown.events)


def test_the_cli_and_the_mcp_tools_take_the_ceiling(monkeypatch):
    import sys

    import pytest

    import aquascope.studio
    from aquascope import cli

    seen: dict = {}

    class Fake:
        def __init__(self, *a, **kw):
            seen.update(kw)
            raise ValueError("stop here")

    monkeypatch.setattr(aquascope.studio, "Studio", Fake)
    monkeypatch.setattr(sys, "argv", ["aquascope", "studio", "q", "--lat", "1", "--lon", "2", "--max-usd", "1.5"])
    with pytest.raises(SystemExit):
        cli.main()
    assert seen["max_usd"] == 1.5
    import inspect

    from aquascope import mcp_server

    for name in ("studio_start", "studio_say", "studio_approve", "studio_follow_up"):
        assert "max_usd" in inspect.signature(getattr(mcp_server, name)).parameters, name


def test_a_figure_is_listed_once_and_the_raw_record_stays_in_the_workbook():
    from aquascope.studio.deliverables import report_md, workbook

    ws = _ran()
    png = Artifact(id="s2_series", kind="figure", name="figures/s2_series.png", data=b"\x89PNG\r\n\x1a\n",
                   media_type="image/png", caption="the series", step="s2", meta={"kind": "series"})
    svg = Artifact(id="s2_series-svg", kind="figure", name="figures/s2_series.svg", data=b"<svg/>",
                   media_type="image/svg+xml", caption="the series", step="s2", meta={"kind": "series"})
    rows = "\n".join(f"1990-01-{d:02d},{10 + d}" for d in range(1, 31))
    series = Artifact(id="tab-s2-series", kind="table", name="tables/s2_series.csv",
                      data=f"datetime,value\n{rows}\n".encode(), media_type="text/csv",
                      caption="The record at Kingston (datetime, value).", step="s2",
                      meta={"columns": ["datetime", "value"], "rows": 30, "name": "series"})
    summary = Artifact(id="tab-s2-summary", kind="table", name="tables/s2_summary.csv",
                       data=b"statistic,value\nmean,12.5\n", media_type="text/csv", caption="Summary.", step="s2",
                       meta={"columns": ["statistic", "value"], "rows": 1, "name": "summary"})
    for a in (png, svg, series, summary):
        ws.add_artifact(a)
    author.author_report(ws, None)
    sec = next(x for x in ws.report["sections"] if x["id"] == "results-s2")
    assert sec["figures"] == ["s2_series"], "the SVG twin is the same figure"
    md = report_md.report_markdown(ws)
    assert md.count("](figures/s2_series.png)") == 1
    assert "1990-01-07" not in md and "sheet `s2_series`" in md and "| mean | 12.5 |" in md
    html = report_md.report_html(ws)
    n_png = sum(1 for a in ws.artifacts if a.kind == "figure" and a.media_type == "image/png")
    assert html.count("data:image/png;base64,") == n_png and "1990-01-07" not in html
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(workbook.workbook_bytes(ws)), read_only=True)
    assert any("series" in name for name in book.sheetnames), book.sheetnames
